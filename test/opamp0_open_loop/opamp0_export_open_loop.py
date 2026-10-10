#!/usr/bin/env python3
"""Export the fully open-loop DC snapshot from Opamp0_open_loop.cir to Excel.

ngspice invokes this script after writing the nine-metric result.
Run it manually to recover Excel from an already completed simulation.

Dependency:
    python3 -m pip install openpyxl

Manual usage:
    python3 opamp0_export_open_loop.py --netlist Opamp0_open_loop.cir
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile


METRICS = (
    ("DC_Gain_dB", "dB"),
    ("UGF_Hz", "Hz"),
    ("Phase_Margin_deg", "deg"),
    ("CMRR_dB", "dB"),
    ("PSRR_Plus_dB", "dB"),
    ("PSRR_Minus_dB", "dB"),
    ("Slew_Rise_V_us", "V/us"),
    ("Slew_Fall_V_us", "V/us"),
    ("Power_Quiescent_uW", "uW"),
)

CAPS = (
    "cgg", "cgs", "cgd", "cgb",
    "cdg", "cdd", "cds", "cdb",
    "csg", "csd", "css", "csb",
    "cbg", "cbd", "cbs", "cbb",
    "capbd", "capbs",
)

CHARGES = ("qg", "qb", "qd", "qs")

CORE = (
    "w", "l", "m",
    "id", "gm", "gds", "gmbs",
    "vth", "vdsat",
)

DC_HEADERS = (
    "Device",
    "MOS_Type",
    "Model",
    "D_Node",
    "G_Node",
    "S_Node",
    "B_Node",
    "W_um",
    "L_um",
    "M",
    "VD_V",
    "VG_V",
    "VS_V",
    "VB_V",
    "VGS_V",
    "VDS_V",
    "VBS_V",
    "Id_model_A",
    "gm_S",
    "gds_S",
    "gmbs_S",
    "Vth_model_V",
    "Vdsat_model_V",
    "gm_over_Id_per_V",
    "gm_over_gds",
    "ro_Ohm",
    "VDS_oriented_V",
    "Vov_est_V",
    "Saturation_margin_V",
)

CAP_HEADERS = (
    ("Device", "MOS_Type", "M")
    + tuple(q + "_F" for q in CAPS)
    + tuple(q + "_C" for q in CHARGES)
)

# 18 个独立设计参数：
# 7 组 W/L + M0/M1/M2 + IBIAS_A。
PARAMETERS = tuple(
    name
    for i in range(1, 8)
    for name in (f"DESVAR_W{i}", f"DESVAR_L{i}")
)
PARAMETERS += (
    tuple(f"DESVAR_M{i}" for i in range(3))
    + ("IBIAS_A",)
)

CONDITIONS = (
    ("SUPPLY_VDD", "V"),
    ("SUPPLY_VSS", "V"),
    ("INPUT_VCM", "V"),
    ("INPUT_VDIFF", "V"),
    ("LOAD_C", "F"),
    ("AC_DENSITY", "points/dec"),
    ("AC_START_HZ", "Hz"),
    ("AC_STOP_HZ", "Hz"),
    ("REF_HZ", "Hz"),
    ("SR_DIFF_LOW", "V"),
    ("SR_DIFF_HIGH", "V"),
)


def finite_or_marker(value: float):
    """Excel 中用字符串保存 NaN/Inf，保留正常数值及负值。"""
    if math.isnan(value):
        return "nan"
    if math.isinf(value):
        return "inf" if value > 0 else "-inf"
    return value


def read_dc_raw(path: Path) -> dict[str, float]:
    """读取包含一个实数 DC 工作点的 ngspice ASCII raw 文件。"""
    lines = path.read_text(
        encoding="utf-8",
        errors="strict",
    ).splitlines()

    header = {}

    try:
        start = lines.index("Variables:")
        end = lines.index("Values:", start)
    except ValueError as exc:
        raise ValueError(
            "Expected an ASCII ngspice raw file "
            "with Variables/Values sections"
        ) from exc

    for line in lines[:start]:
        if ":" in line:
            key, value = line.split(":", 1)
            header[key.strip().lower()] = value.strip()

    if header.get("plotname", "").lower() != "operating point":
        raise ValueError(
            "The raw file is not a DC Operating Point plot"
        )

    if header.get("flags", "").lower().split() != ["real"]:
        raise ValueError(
            "Expected real DC values, not a complex AC plot"
        )

    if int(header.get("no. points", "0")) != 1:
        raise ValueError("Expected exactly one DC point")

    count = int(header["no. variables"])
    variables = []

    for line in lines[start + 1:end]:
        fields = line.split()

        if fields:
            if (
                len(fields) < 3
                or int(fields[0]) != len(variables)
            ):
                raise ValueError(
                    "Malformed raw-file variable table"
                )

            variables.append(fields[1].lower())

    if (
        len(variables) != count
        or len(set(variables)) != count
    ):
        raise ValueError(
            "Incomplete or duplicated variable table"
        )

    data = [
        line.split()
        for line in lines[end + 1:]
        if line.strip()
    ]

    if (
        len(data) != count
        or len(data[0]) != 2
        or data[0][0] != "0"
    ):
        raise ValueError("Incomplete raw-file value table")

    values = [float(data[0][1])]

    for fields in data[1:]:
        if len(fields) != 1:
            raise ValueError("Malformed raw-file value")

        values.append(float(fields[0]))

    if not all(math.isfinite(value) for value in values):
        raise ValueError(
            "Non-finite DC state; refusing to present it "
            "as a completed snapshot"
        )

    return dict(zip(variables, values))


def parse_devices(netlist: str):
    """从 DUT 子电路读取器件名、型号和端口连接。"""
    match = re.search(
        r"(?ims)^\.subckt\s+DUT\s+VINP\s+VINN\s+VOUT"
        r"\s+VDD\s+VSS\s*\n(.*?)^\.ends\s+DUT\b",
        netlist,
    )

    if not match:
        raise ValueError(
            "Cannot find the canonical DUT interface "
            "in the supplied netlist"
        )

    devices = []

    for line in match.group(1).splitlines():
        fields = line.split()

        if (
            len(fields) >= 6
            and fields[0].lower().startswith(("xpm", "xnm"))
        ):
            name, drain, gate, source, bulk, model = fields[:6]
            model = model.lower()

            if model not in (
                "sky130_fd_pr__pfet_01v8",
                "sky130_fd_pr__nfet_01v8",
            ):
                raise ValueError(
                    f"Unsupported MOS model: {model}"
                )

            devices.append(
                (
                    name.upper(),
                    "PMOS" if "pfet" in model else "NMOS",
                    model,
                    drain.upper(),
                    gate.upper(),
                    source.upper(),
                    bulk.upper(),
                )
            )

    if (
        len(devices) != 17
        or len({device[0] for device in devices}) != 17
    ):
        raise ValueError(
            "This testbench must contain "
            "the expected 17 unique MOS devices"
        )

    return devices


def read_metrics(path: Path):
    """读取九项指标，保留 NaN、Inf 和有效负值。"""
    with path.open(
        newline="",
        encoding="utf-8",
    ) as handle:
        rows = list(csv.reader(handle))

    if (
        len(rows) != 2
        or rows[0] != [metric[0] for metric in METRICS]
        or len(rows[1]) != len(METRICS)
    ):
        raise ValueError(
            "Metrics TXT must have the exact canonical "
            "header and one nine-column row"
        )

    return [
        [name, finite_or_marker(float(cell)), unit]
        for (name, unit), cell in zip(METRICS, rows[1])
    ]


def dc_formulas(row: int):
    """工作点表中派生量的 Excel 公式。"""
    # 原生 Id/gm/gds 已包含 M，不再额外乘 M。
    return {
        "gm_over_Id_per_V": (
            f'=IF(ABS(R{row})>0,'
            f'ABS(S{row})/ABS(R{row}),"nan")'
        ),
        "gm_over_gds": (
            f'=IF(ABS(T{row})>0,'
            f'ABS(S{row})/ABS(T{row}),"nan")'
        ),
        "ro_Ohm": (
            f'=IF(ABS(T{row})>0,'
            f'1/ABS(T{row}),"nan")'
        ),
        "VDS_oriented_V": (
            f'=IF(B{row}="PMOS",-P{row},P{row})'
        ),
        "Vov_est_V": (
            f'=IF(B{row}="PMOS",-O{row},O{row})'
            f'-ABS(V{row})'
        ),
        "Saturation_margin_V": (
            f"=AA{row}-ABS(W{row})"
        ),
    }


def prepare_payload(
    raw: dict[str, float],
    netlist_text: str,
    metric_rows,
):
    """将 raw、网表和指标文件整理成统一导出结构。"""

    def required(name):
        try:
            return raw[name.lower()]
        except KeyError as exc:
            raise ValueError(
                f"Missing required DC vector {name}; "
                "the netlist/export schema may be stale"
            ) from exc

    def voltage(node):
        port_map = {
            "VINP": "inp_cm",
            "VINN": "inn_cm",
            "VOUT": "out_cm",
            "VDD": "vdd_q",
            "VSS": "vss_q",
        }

        if node in ("0", "GND"):
            return 0.0

        resolved_node = port_map.get(
            node,
            "xdut_cm." + node.lower(),
        )
        return required(f"v({resolved_node})")

    dc_rows = []
    cap_rows = []

    for device in parse_devices(netlist_text):
        name, kind, model, drain, gate, source, bulk = device

        fields = {
            quantity: required(
                f"op_{name.lower()}_{quantity}"
            )
            for quantity in CORE + CAPS + CHARGES
        }

        vd, vg, vs, vb = (
            voltage(node)
            for node in (drain, gate, source, bulk)
        )

        vgs = vg - vs
        vds = vd - vs
        vbs = vb - vs

        polarity = -1 if kind == "PMOS" else 1
        oriented = polarity * vds

        ratio = (
            abs(fields["gm"]) / abs(fields["id"])
            if fields["id"]
            else math.nan
        )

        intrinsic = (
            abs(fields["gm"]) / abs(fields["gds"])
            if fields["gds"]
            else math.nan
        )

        ro = (
            1 / abs(fields["gds"])
            if fields["gds"]
            else math.nan
        )

        row = [
            name,
            kind,
            model,
            drain,
            gate,
            source,
            bulk,
            fields["w"] * 1e6,
            fields["l"] * 1e6,
            fields["m"],
            vd,
            vg,
            vs,
            vb,
            vgs,
            vds,
            vbs,
            fields["id"],
            fields["gm"],
            fields["gds"],
            fields["gmbs"],
            fields["vth"],
            fields["vdsat"],
            ratio,
            intrinsic,
            ro,
            oriented,
            polarity * vgs - abs(fields["vth"]),
            oriented - abs(fields["vdsat"]),
        ]

        dc_rows.append(
            [
                finite_or_marker(value)
                if isinstance(value, float)
                else value
                for value in row
            ]
        )

        cap_rows.append(
            [name, kind, fields["m"]]
            + [
                fields[quantity]
                for quantity in CAPS + CHARGES
            ]
        )

    setup_rows = [
        [
            "DC snapshot",
            "Instance",
            "XDUT_CM",
            "",
            "Open-loop DUT; independent DC inputs; "
            "same instance as quiescent power",
        ],
        [
            "DC snapshot",
            "VINP_actual",
            required("v(inp_cm)"),
            "V",
            "Actual node voltage, not the requested target",
        ],
        [
            "DC snapshot",
            "VINN_actual",
            required("v(inn_cm)"),
            "V",
            "Independently biased inverting input; "
            "no output feedback",
        ],
        [
            "DC snapshot",
            "VOUT_actual",
            required("v(out_cm)"),
            "V",
            "Actual solved DC output",
        ],
        [
            "DC snapshot",
            "Output_minus_VINP",
            required("output_minus_vinp_v"),
            "V",
            "VOUT minus VINP; not a follower tracking error",
        ],
        [
            "DC snapshot",
            "Input_differential",
            required("input_differential_v"),
            "V",
            "Actual VINP minus VINN",
        ],
        [
            "DC snapshot",
            "Output_headroom_low",
            required("output_headroom_low_v"),
            "V",
            "VOUT minus VSS; inspect bias "
            "before interpreting gain",
        ],
        [
            "DC snapshot",
            "Output_headroom_high",
            required("output_headroom_high_v"),
            "V",
            "VDD minus VOUT; inspect bias "
            "before interpreting gain",
        ],
        [
            "DC snapshot",
            "DM_CM_bias_difference",
            required("dm_cm_bias_difference_v"),
            "V",
            "OUT_DM minus OUT_CM",
        ],
        [
            "DC snapshot",
            "Power_Quiescent",
            required("power_uw"),
            "uW",
            "One DUT instance only",
        ],
    ]

    for name, unit in CONDITIONS:
        setup_rows.append(
            [
                "Test condition",
                name,
                required("op_input_" + name.lower()),
                unit,
                "",
            ]
        )

    temp = re.search(
        r"(?im)^\s*\.temp\s+([-+\d.eE]+)",
        netlist_text,
    )

    if temp:
        setup_rows.append(
            [
                "Test condition",
                "TEMP",
                float(temp.group(1)),
                "degC",
                "",
            ]
        )

    pdk = re.search(
        r'(?im)^\s*\.lib\s+"([^"]+)"\s+(\S+)',
        netlist_text,
    )

    if pdk:
        setup_rows.extend(
            [
                [
                    "Test condition",
                    "PDK_PATH",
                    pdk.group(1),
                    "",
                    "Model library used for this run",
                ],
                [
                    "Test condition",
                    "CORNER",
                    pdk.group(2),
                    "",
                    "",
                ],
            ]
        )

    # 参数数值从本次仿真的 raw 快照读取，
    # 不使用 Python 中硬编码的示例值。
    for name in PARAMETERS:
        if name.startswith(("DESVAR_W", "DESVAR_L")):
            unit = "um"
        elif name == "IBIAS_A":
            unit = "A"
        else:
            unit = "count"

        setup_rows.append(
            [
                "Design parameter",
                name,
                required("op_input_" + name.lower()),
                unit,
                "Resolved numeric value",
            ]
        )

    setup_rows.extend(
        [
            [
                "Definition",
                "VGS_V / VDS_V / VBS_V",
                "Terminal node differences",
                "V",
                "Signed; PMOS VGS/VDS normally negative",
            ],
            [
                "Definition",
                "Id_model_A / gm_S / gds_S",
                "Native ngspice BSIM values",
                "",
                "Already include M; no extra multiplication",
            ],
            [
                "Definition",
                "Capacitances",
                "Native BSIM charge-derivative matrix",
                "F",
                "Preserve signed off-diagonal terms; "
                "not all entries are positive lumped capacitors",
            ],
            [
                "Definition",
                "Saturation_margin_V",
                "VDS_oriented minus abs(Vdsat_model)",
                "V",
                "Diagnostic only; not an unconditional "
                "operating-region classifier",
            ],
            [
                "Definition",
                "Vov_est_V",
                "Oriented VGS minus abs(Vth_model)",
                "V",
                "Diagnostic only; negative overdrive "
                "can occur in weak inversion",
            ],
            [
                "Definition",
                "DC_Gain_dB / UGF_Hz",
                "Differential open-loop gain Ad "
                "at fixed common mode",
                "",
                "UGF is the first falling crossing "
                "of abs(Ad)=1",
            ],
            [
                "Definition",
                "CMRR_dB",
                "20*log10(abs(Ad/Acm))",
                "dB",
                "Direct open-loop common-mode injection",
            ],
            [
                "Definition",
                "PSRR_Plus_dB / PSRR_Minus_dB",
                "20*log10(abs(Ad/Asupply))",
                "dB",
                "Input-referred rejection; "
                "independent inputs are AC grounded",
            ],
            [
                "Definition",
                "Phase_Margin_deg",
                "Hypothetical unity-feedback estimate "
                "at this open-loop bias",
                "deg",
                "T=Ad-Acm/2=-gn; PM at T unity crossing. "
                "Closing feedback may change the DC bias. "
                "Not a stability proof.",
            ],
            [
                "Definition",
                "Slew_Rise_V_us / Slew_Fall_V_us",
                "Open-loop output transition rates",
                "V/us",
                "10%-90% of window-averaged swing. "
                "Inspect transient plateaus. "
                "Not follower slew rates.",
            ],
            [
                "Definition",
                "Missing metric",
                "nan",
                "",
                "No fabricated zero; "
                "valid negative metrics are retained",
            ],
            [
                "Definition",
                "Transient initial point",
                "Separate from the quiet DC snapshot",
                "",
                "Differential pulse starts at "
                "INPUT_VDIFF+SR_DIFF_LOW; "
                "fixed input common mode",
            ],
        ]
    )

    return {
        "schema_version": 1,
        "dc_headers": list(DC_HEADERS),
        "dc_rows": dc_rows,
        "cap_headers": list(CAP_HEADERS),
        "cap_rows": cap_rows,
        "metrics": metric_rows,
        "setup_rows": setup_rows,
        "dc_notes": (
            "DC snapshot: XDUT_CM fully open-loop; "
            "independently biased inputs. "
            "Signed terminal voltages; "
            "model currents already include M."
        ),
        "cap_notes": (
            "Native BSIM capacitance/charge matrix. "
            "Preserve signs of cross terms."
        ),
    }


def atomic_csv(path: Path, header, rows):
    """先写临时文件，成功后替换目标 CSV。"""
    path.parent.mkdir(parents=True, exist_ok=True)

    fd, temporary = tempfile.mkstemp(
        prefix=path.name + ".",
        suffix=".tmp",
        dir=path.parent,
    )

    try:
        with os.fdopen(
            fd,
            "w",
            newline="",
            encoding="utf-8-sig",
        ) as handle:
            writer = csv.writer(handle)
            writer.writerow(header)
            writer.writerows(rows)

        os.replace(temporary, path)

    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def write_wide_csv(path: Path, payload):
    """每个设计点一行，所有 MOS 工作点字段展开为列。"""
    header = ["DC_Instance"]
    values = ["XDUT_CM"]

    for row, cap in zip(
        payload["dc_rows"],
        payload["cap_rows"],
    ):
        for field, value in zip(
            payload["dc_headers"][7:],
            row[7:],
        ):
            header.append(row[0] + "__" + field)
            values.append(value)

        for field, value in zip(
            payload["cap_headers"][3:],
            cap[3:],
        ):
            header.append(row[0] + "__" + field)
            values.append(value)

    atomic_csv(path, header, [values])


def write_user_workbook(path: Path, payload):
    """生成工作点、电容、指标及参数四个工作表。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.workbook.properties import CalcProperties

    except ImportError as exc:
        raise RuntimeError(
            "Install the Excel dependency first: "
            "python3 -m pip install openpyxl"
        ) from exc

    workbook = Workbook()
    workbook.remove(workbook.active)

    workbook.calculation = CalcProperties(
        calcId=124519,
        fullCalcOnLoad=True,
    )

    contents = (
        (
            "DC_Operating_Point",
            payload["dc_headers"],
            payload["dc_rows"],
            payload["dc_notes"],
        ),
        (
            "MOS_Capacitances",
            payload["cap_headers"],
            payload["cap_rows"],
            payload["cap_notes"],
        ),
        (
            "Metrics",
            ["Metric", "Value", "Unit"],
            payload["metrics"],
            "",
        ),
        (
            "Setup",
            ["Category", "Parameter", "Value", "Unit", "Notes"],
            payload["setup_rows"],
            "",
        ),
    )

    for name, headers, rows, note in contents:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False

        sheet.append(headers)

        for row in rows:
            sheet.append(row)

        if name == "DC_Operating_Point":
            sheet.freeze_panes = "H2"
        elif name == "MOS_Capacitances":
            sheet.freeze_panes = "D2"
        else:
            sheet.freeze_panes = "A2"

        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.fill = PatternFill(
                "solid",
                fgColor="253B55",
            )
            cell.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color="FFFFFF",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sheet.row_dimensions[1].height = 34

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(
                    name="Arial",
                    size=10,
                    color="172B4D",
                )

                cell.alignment = Alignment(
                    horizontal=(
                        "right"
                        if isinstance(cell.value, (int, float))
                        else "left"
                    ),
                    vertical="center",
                )

                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000E+00"

            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="F3F6FA",
                    )

        for col_index, header in enumerate(headers, 1):
            width = min(
                max(len(str(header)) + 2, 16),
                32,
            )

            if (
                name == "DC_Operating_Point"
                and col_index == 3
            ):
                width = 32

            if name == "Setup":
                width = [22, 30, 68, 14, 100][col_index - 1]

            if name == "Metrics":
                width = [29, 22, 14][col_index - 1]

            sheet.column_dimensions[
                get_column_letter(col_index)
            ].width = width

        if note:
            sheet.cell(len(rows) + 3, 1, note)

        if name == "DC_Operating_Point":
            for row_index in range(2, len(rows) + 2):
                for field, formula in dc_formulas(
                    row_index
                ).items():
                    column = headers.index(field) + 1
                    cell = sheet.cell(
                        row_index,
                        column,
                        formula,
                    )
                    cell.number_format = "0.000000E+00"

            for column in (8, 9):
                for row in sheet.iter_rows(
                    min_row=2,
                    max_row=len(rows) + 1,
                    min_col=column,
                    max_col=column,
                ):
                    row[0].number_format = "0.000"

            for row in sheet.iter_rows(
                min_row=2,
                max_row=len(rows) + 1,
                min_col=10,
                max_col=10,
            ):
                row[0].number_format = "0"

    path.parent.mkdir(parents=True, exist_ok=True)

    fd, temporary = tempfile.mkstemp(
        prefix=path.stem + ".",
        suffix=".xlsx",
        dir=path.parent,
    )
    os.close(fd)

    try:
        workbook.save(temporary)
        os.replace(temporary, path)

    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main(argv=None):
    parser = argparse.ArgumentParser(
        description=__doc__,
    )

    parser.add_argument(
        "--netlist",
        type=Path,
        default=Path("Opamp0_open_loop.cir"),
    )
    parser.add_argument(
        "--raw",
        type=Path,
        default=Path("opamp0_openloop_dc.raw"),
    )
    parser.add_argument(
        "--metrics",
        type=Path,
        default=Path("opamp0_openloop_metrics.txt"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("opamp0_openloop_results.xlsx"),
    )
    parser.add_argument(
        "--wide-csv",
        type=Path,
        default=Path("opamp0_openloop_dc.csv"),
    )
    parser.add_argument(
        "--json-output",
        type=Path,
        help="Optional numeric intermediate for analysis or CI",
    )
    parser.add_argument(
        "--json-only",
        action="store_true",
        help=(
            "Validate/export the numeric intermediate "
            "without XLSX"
        ),
    )

    args = parser.parse_args(argv)

    try:
        payload = prepare_payload(
            read_dc_raw(args.raw),
            args.netlist.read_text(encoding="utf-8"),
            read_metrics(args.metrics),
        )

        write_wide_csv(args.wide_csv, payload)

        if args.json_output or args.json_only:
            target = (
                args.json_output
                or args.output.with_suffix(".json")
            )

            target.parent.mkdir(
                parents=True,
                exist_ok=True,
            )

            target.write_text(
                json.dumps(
                    payload,
                    ensure_ascii=False,
                    indent=2,
                    allow_nan=False,
                ),
                encoding="utf-8",
            )

        if not args.json_only:
            write_user_workbook(args.output, payload)

            print(
                f"Excel saved: {args.output} "
                f"(17 MOS devices, 9 metrics, "
                f"{len(PARAMETERS)} design parameters)",
                flush=True,
            )

        print(
            f"DC wide CSV saved: {args.wide_csv}",
            flush=True,
        )

        return 0

    except (OSError, ValueError, RuntimeError) as exc:
        print(
            f"Excel/DC export failed: {exc}",
            file=sys.stderr,
            flush=True,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())