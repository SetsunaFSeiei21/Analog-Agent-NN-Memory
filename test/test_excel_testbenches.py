"""Validate the independent Excel benches without claiming electrical validation.

Run from the repository root:
    python3 -m unittest discover -s test -p 'test_excel_testbenches.py' -v

The ASCII raw fixtures below are synthetic export inputs, not SPICE results.
Install openpyxl to include XLSX validation.
"""
from __future__ import annotations

import ast
import csv
import importlib.util
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "analog-agent/Sample_Optimizer_Circuit"
COUNTS = {
    "5t_ota": (6, 8),
    "Opamp1": (24, 22),
    "two_stage_opamp_otaf": (8, 13),
    "two_stage_folded_opamp": (22, 23),
}
QUANTITIES = (
    "w l m id gm gds gmbs vth vdsat "
    "cgg cgs cgd cgb cdg cdd cds cdb csg csd css csb cbg cbd cbs cbb "
    "capbd capbs qg qb qd qs"
).split()
METRICS = (
    "DC_Gain_dB UGF_Hz Phase_Margin_deg CMRR_dB PSRR_Plus_dB "
    "PSRR_Minus_dB Slew_Rise_V_us Slew_Fall_V_us Power_Quiescent_uW"
).split()


def dut(text):
    return re.search(r"(?ims)^\.subckt DUT\b.*?^\.ends DUT\b", text).group()


def active_lines(text):
    return [" ".join(line.lower().split()) for line in text.splitlines()
            if line.strip() and not line.lstrip().startswith("*")]


def assignments(text):
    result = {}
    for line in text.splitlines():
        if line.lower().startswith(".param "):
            for name, value in re.findall(r"(\w+)\s*=\s*(\S+)", line[7:]):
                result[name] = value
    return result


def expression_value(expression, parameters):
    """Evaluate only the arithmetic used by these checked-in sizing expressions."""
    def visit(node):
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.Name):
            return parameters[node.id]
        if isinstance(node, ast.BinOp):
            left, right = visit(node.left), visit(node.right)
            if isinstance(node.op, ast.Mult):
                return left * right
            if isinstance(node.op, ast.Add):
                return left + right
        raise ValueError(f"Unsupported test expression: {expression}")
    return visit(ast.parse(expression.strip("{}"), mode="eval").body)


def fixture(circuit, exporter, text, changed=False):
    # Obtain design values from the independent source files, not the exporter.
    source_parameters = assignments((SOURCES / circuit / f"{circuit}_params.sp").read_text())
    parameters = {name: float(value) for name, value in source_parameters.items()}
    if changed:
        for name in parameters:
            unit = exporter.PARAMETER_UNITS[name]
            if unit == "count":
                parameters[name] += 2
            else:
                parameters[name] *= {"um": 1.25, "A": 2, "Ohm": 3, "F": 4}[unit]

    raw = {"op_input_" + name.lower(): value for name, value in parameters.items()}
    conditions = dict(SUPPLY_VDD=1.8, SUPPLY_VSS=0.0, INPUT_VCM=0.72, LOAD_C=1e-12,
                      AC_DENSITY=100, AC_START_HZ=0.1, AC_STOP_HZ=50e9, REF_HZ=1,
                      SR_LOW=0.7, SR_HIGH=1.1)
    raw.update(("op_input_" + name.lower(), value) for name, value in conditions.items())
    ports = {"VINP": "inp_cm", "VINN": "inn_cm", "VOUT": "out_cm", "VDD": "vdd_q", "VSS": "vss_q"}
    raw.update({"v(inp_cm)": 0.72, "v(inn_cm)": 0.718, "v(out_cm)": 0.718,
                "v(vdd_q)": 1.8, "v(vss_q)": 0.0, "power_uw": 72.0,
                "follower_error_v": -0.002, "dm_cm_bias_difference_v": 0.0})
    for index, line in enumerate(dut(text).splitlines()):
        if not line.startswith("X"):
            continue
        name, drain, gate, source, bulk, model, *_ = line.split()
        for node in (drain, gate, source, bulk):
            if node not in ports:
                raw.setdefault("v(xdut_cm." + node.lower() + ")", 0.1 * (1 + index % 12))
        sizing = dict((key.lower(), value) for key, value in
                      re.findall(r"(\w+)\s*=\s*(\{[^}]+\}|\S+)", line))
        m = expression_value(sizing["m"], parameters)
        sign = -1 if "pfet" in model else 1
        state = {"w": expression_value(sizing["w"], parameters) * 1e-6,
                 "l": expression_value(sizing["l"], parameters) * 1e-6,
                 "m": m, "id": sign * 2e-6 * m, "gm": 20e-6 * m,
                 "gds": 0.3e-6 * m, "gmbs": 1e-6 * m,
                 "vth": sign * 0.38, "vdsat": sign * 0.1}
        for i, quantity in enumerate(QUANTITIES[9:]):
            state[quantity] = -(i + 1) * 1e-15 * m
        raw.update((f"op_{name.lower()}_{quantity}", value) for quantity, value in state.items())
    metrics = [[name, value, unit] for (name, unit), value in zip(
        exporter.METRICS, (-2.0, "nan", -10.0, "inf", 40.0, 45.0, "nan", 3.0, 72.0))]
    return raw, metrics, parameters


def write_raw(path, raw):
    lines = ["Title: Synthetic export fixture - NOT a simulated circuit",
             "Plotname: Operating Point", "Flags: real",
             f"No. Variables: {len(raw)}", "No. Points: 1", "Variables:"]
    lines += [f"\t{i}\t{name}\tvoltage" for i, name in enumerate(raw)]
    lines += ["Values:"]
    lines += [("0\t" if i == 0 else "\t") + format(value, ".17e")
              for i, value in enumerate(raw.values())]
    path.write_text("\n".join(lines) + "\n")


class ExcelBenchesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.packages = {}
        for circuit in COUNTS:
            folder = ROOT / "test" / (circuit.lower() + "_excel")
            script = folder / (circuit.lower() + "_export_excel.py")
            spec = importlib.util.spec_from_file_location(circuit + "_exporter", script)
            exporter = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(exporter)
            text = (folder / (circuit + "_all_metrics.cir")).read_text()
            cls.packages[circuit] = folder, exporter, text

    def test_every_other_source_circuit_has_an_independent_package(self):
        found = {p.name for p in SOURCES.iterdir() if (p / (p.name + ".sp")).is_file()}
        self.assertEqual(found - {"Opamp0"}, set(COUNTS))
        for folder, exporter, _ in self.packages.values():
            self.assertEqual((folder / ".spiceinit").read_text(), "set ngbehavior=hs\n")
            self.assertEqual([name for name, _ in exporter.METRICS], METRICS)
            self.assertEqual(set(exporter.CORE + exporter.CAPS + exporter.CHARGES), set(QUANTITIES))

    def test_dut_connectivity_expressions_and_source_defaults_are_preserved(self):
        for circuit, (_, exporter, text) in self.packages.items():
            with self.subTest(circuit=circuit):
                source = (SOURCES / circuit / f"{circuit}.sp").read_text()
                self.assertEqual(active_lines(dut(text).replace("_01v8_lvt", "_01v8")),
                                 active_lines(dut(source).replace("_01v8_lvt", "_01v8")))
                defaults = assignments((SOURCES / circuit / f"{circuit}_params.sp").read_text())
                actual = assignments(text)
                self.assertEqual({name: actual[name] for name in defaults}, defaults)
                self.assertEqual(set(exporter.PARAMETERS), set(defaults))
                self.assertEqual(len(exporter.PARAMETERS), COUNTS[circuit][1])
                self.assertEqual(actual["INPUT_VCM"], "0.72")
                self.assertNotRegex(text, r"(?im)^\.include\b")
                devices = exporter.parse_devices(text)
                self.assertEqual(len(devices), COUNTS[circuit][0])
                self.assertTrue(all(row[2].endswith("_01v8_lvt") for row in devices))

    def test_every_dc_query_targets_the_actual_model_and_all_inputs_are_exported(self):
        for circuit, (_, exporter, text) in self.packages.items():
            with self.subTest(circuit=circuit):
                expected = {}
                for name, _, model, *_ in exporter.parse_devices(text):
                    for quantity in QUANTITIES:
                        expected[f"op_{name.lower()}_{quantity}"] = (
                            f"@m.xdut_cm.{name.lower()}.m{model}[{quantity}]"
                        )
                inputs = list(exporter.PARAMETERS) + [name for name, _ in exporter.CONDITIONS]
                expected.update((f"op_input_{name.lower()}", f"const.export_{name.lower()}") for name in inputs)
                actual = re.findall(r"(?m)^let (op_\w+)=(.+)$", text)
                self.assertEqual(len(actual), len(expected))
                self.assertEqual(dict(actual), expected)
                aliases = re.findall(r"(?m)^\.csparam export_(\w+)=\{(\w+)\}$", text)
                self.assertEqual(aliases, [(name.lower(), name) for name in inputs])
                self.assertLess(text.index("\nop\n"), text.index("let op_input_"))
                self.assertLess(text.index("let op_input_"), text.index("\nwrite "))
                self.assertLess(text.index("\nwrite "), text.index("\nac dec "))

    def test_metric_algorithm_matches_opamp0(self):
        template = (ROOT / "test/opamp0_excel/Opamp0_all_metrics.cir").read_text()
        start, stop = "* One AC sweep supplies", "* Python creates the Excel"
        reference = template[template.index(start):template.index(stop)]
        for circuit, (_, _, text) in self.packages.items():
            with self.subTest(circuit=circuit):
                self.assertEqual(text[text.index(start):text.index(stop)], reference)

    def test_exported_dimensions_multipliers_and_compensation_follow_changed_inputs(self):
        for circuit, (_, exporter, text) in self.packages.items():
            for changed in (False, True):
                with self.subTest(circuit=circuit, changed=changed):
                    raw, metrics, parameters = fixture(circuit, exporter, text, changed)
                    payload = exporter.prepare_payload(raw, text, metrics)
                    setup = {row[1]: row[2:4] for row in payload["setup_rows"] if row[0] == "Design parameter"}
                    self.assertEqual(set(setup), set(parameters))
                    for name, value in parameters.items():
                        self.assertEqual(setup[name], [value, exporter.PARAMETER_UNITS[name]])
                    for values, caps in zip(payload["dc_rows"], payload["cap_rows"]):
                        row = dict(zip(exporter.DC_HEADERS, values))
                        prefix = "op_" + row["Device"].lower() + "_"
                        self.assertAlmostEqual(row["W_um"], raw[prefix + "w"] * 1e6)
                        self.assertAlmostEqual(row["L_um"], raw[prefix + "l"] * 1e6)
                        self.assertEqual(row["M"], raw[prefix + "m"])
                        self.assertEqual(row["Id_model_A"], raw[prefix + "id"])
                        self.assertEqual(row["gm_S"], raw[prefix + "gm"])
                        self.assertAlmostEqual(row["gm_over_Id_per_V"], 10.0)
                        self.assertEqual(caps[3], raw[prefix + "cgg"])
                        self.assertAlmostEqual(row["VBS_V"], row["VB_V"] - row["VS_V"])
                        self.assertAlmostEqual(row["VDS_V"], row["VD_V"] - row["VS_V"])
                    if "RNULL_OHM" in setup:
                        self.assertEqual(setup["RNULL_OHM"][1], "Ohm")
                        self.assertEqual(setup["CCOMP_F"][1], "F")
                    if circuit == "Opamp1":
                        kinds = {row[0]: row[1] for row in payload["dc_rows"]}
                        self.assertEqual((kinds["XM10"], kinds["XM11"]), ("PMOS", "NMOS"))

    def test_missing_vectors_unsupported_models_and_wrong_device_inventory_fail(self):
        for circuit, (_, exporter, text) in self.packages.items():
            with self.subTest(circuit=circuit):
                raw, metrics, _ = fixture(circuit, exporter, text)
                device = exporter.EXPECTED_DEVICES[0]
                del raw[f"op_{device.lower()}_gm"]
                with self.assertRaisesRegex(ValueError, "Missing required DC vector"):
                    exporter.prepare_payload(raw, text, metrics)
                line = next(line for line in dut(text).splitlines() if line.startswith(device + " "))
                for bad in (text.replace(line, "", 1), text.replace(line, line + "\n" + line, 1),
                            text.replace(line, line.replace(device + " ", "XUNKNOWN ", 1), 1)):
                    with self.assertRaisesRegex(ValueError, "unique MOS devices"):
                        exporter.parse_devices(bad)
                unsupported = text.replace("sky130_fd_pr__nfet_01v8_lvt", "unknown_mos", 1)
                with self.assertRaisesRegex(ValueError, "Unsupported MOS model"):
                    exporter.parse_devices(unsupported)

    def run_export(self, circuit, with_excel):
        folder, exporter, text = self.packages[circuit]
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            prefix = circuit.lower()
            raw, metric_rows, _ = fixture(circuit, exporter, text, changed=True)
            write_raw(output / f"{prefix}_dc_operating_point.raw", raw)
            self.assertEqual(exporter.read_dc_raw(output / f"{prefix}_dc_operating_point.raw"), raw)
            (output / f"{circuit}_all_metrics.cir").write_text(text)
            with (output / f"{prefix}_metrics.txt").open("w", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(METRICS)
                writer.writerow([row[1] for row in metric_rows])
            command = [sys.executable, str(folder / f"{prefix}_export_excel.py")]
            command += ["--json-output", "payload.json"] if with_excel else ["--json-only"]
            result = subprocess.run(command, cwd=output, capture_output=True, text=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            payload = json.loads((output / ("payload.json" if with_excel else f"{prefix}_results.json")).read_text())
            self.assertEqual(len(payload["dc_rows"]), COUNTS[circuit][0])
            self.assertEqual(payload["metrics"], metric_rows)
            with (output / f"{prefix}_dc_operating_point.csv").open(encoding="utf-8-sig", newline="") as handle:
                wide = list(csv.reader(handle))
            self.assertEqual(len(wide), 2)
            self.assertEqual(len(wide[0]), len(wide[1]))
            for values in payload["dc_rows"]:
                self.assertIn(values[0] + "__W_um", wide[0])
                self.assertIn(values[0] + "__gm_S", wide[0])
                self.assertIn(values[0] + "__qs_C", wide[0])
            if with_excel:
                import openpyxl
                book = openpyxl.load_workbook(output / f"{prefix}_results.xlsx")
                self.assertEqual(book.sheetnames, ["DC_Operating_Point", "MOS_Capacitances", "Metrics", "Setup"])
                sheet = book["DC_Operating_Point"]
                self.assertEqual([sheet.cell(i + 2, 1).value for i in range(COUNTS[circuit][0])],
                                 list(exporter.EXPECTED_DEVICES))
                for i, values in enumerate(payload["dc_rows"], 2):
                    for j, value in enumerate(values[:23], 1):
                        actual = sheet.cell(i, j).value
                        if isinstance(value, (int, float)):
                            self.assertAlmostEqual(actual, value)
                        else:
                            self.assertEqual(actual, value)
                    for field, formula in exporter.dc_formulas(i).items():
                        self.assertEqual(sheet.cell(i, exporter.DC_HEADERS.index(field) + 1).value, formula)
                self.assertEqual(book["Metrics"].cell(3, 2).value, "nan")
                self.assertEqual(book["Metrics"].cell(4, 2).value, -10)
                self.assertTrue(book.calculation.fullCalcOnLoad)
                book.close()

    def test_standalone_cli_reads_raw_and_writes_numeric_json_and_wide_csv(self):
        for circuit in self.packages:
            with self.subTest(circuit=circuit):
                self.run_export(circuit, with_excel=False)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is needed for XLSX checks")
    def test_standalone_cli_writes_four_sheet_workbooks(self):
        for circuit in self.packages:
            with self.subTest(circuit=circuit):
                self.run_export(circuit, with_excel=True)


if __name__ == "__main__":
    unittest.main()
